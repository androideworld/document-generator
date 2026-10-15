"""Сохранение и загрузка настроек.

Режимы работы:
    Портабельный — config.json в папке рядом с .exe
    Классический — config.json в ~/.config/ГенераторДокументов/ (Linux/macOS)
                   или %APPDATA%\\ГенераторДокументов\\ (Windows)

Портабельный режим активируется автоматически, если рядом с .exe
существует файл config.json. Если нет — используется классический.
"""

import json
import os
import shutil
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional


APP_DIR_NAME = "ГенераторДокументов"
CONFIG_FILENAME = "config.json"

# --- Дефолтные значения ---

DEFAULT_MARKERS: List[Dict[str, Any]] = [
    {"marker": "ClientName",   "column": "ФИО",           "format": "text",   "required": True},
    {"marker": "Amount",       "column": "Сумма",         "format": "amount", "required": True},
    {"marker": "ContractDate", "column": "Дата",          "format": "date",   "required": True},
    {"marker": "ContractNum",  "column": "НомерДоговора", "format": "text",   "required": True},
    {"marker": "Address",      "column": "Адрес",         "format": "text",   "required": False},
    {"marker": "ID",           "column": "ID",            "format": "text",   "required": False},
    {"marker": "Phone",        "column": "Телефон",       "format": "text",   "required": False},
]

DEFAULT_SETTINGS: Dict[str, Any] = {
    "folder_pattern": "{year}/{month}",
    "client_pattern": "{short_fio}_{date}_{contract}",
    "on_missing_marker": "highlight_and_warn",   # highlight_and_warn | highlight | keep
    "on_missing_column": "warn",                 # warn | silent
    "on_empty_value": "empty_string",            # empty_string | keep_marker | placeholder
    "empty_string_spaces": 5,                    # количество пробелов для empty_string
    "placeholder_underscores": 10,               # количество прочерков для placeholder
    "log_mode": "per_run",                       # per_run | single_file
    "ui_theme": "default",
}


# ---------------------------------------------------------------------------
# Определение путей
# ---------------------------------------------------------------------------

def get_install_dir() -> Path:
    """Папка, где лежит утилита.

    - В PyInstaller-сборке — папка, где находится .exe.
    - При запуске из исходников — корень проекта (родитель core/).
    """
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent.parent


def get_classic_config_path() -> Path:
    """Классический путь к config.json (в домашнем каталоге)."""
    if sys.platform == "win32":
        base = os.environ.get("APPDATA") or str(Path.home() / "AppData" / "Roaming")
        return Path(base) / APP_DIR_NAME / CONFIG_FILENAME
    xdg = os.environ.get("XDG_CONFIG_HOME")
    if xdg:
        return Path(xdg) / APP_DIR_NAME / CONFIG_FILENAME
    return Path.home() / ".config" / APP_DIR_NAME / CONFIG_FILENAME


def get_portable_config_path() -> Path:
    """Путь к config.json в папке установки (портабельный режим)."""
    return get_install_dir() / CONFIG_FILENAME


def is_portable() -> bool:
    """True, если config.json уже лежит рядом с .exe."""
    return get_portable_config_path().exists()


def get_config_path() -> Path:
    """Активный путь к config.json.

    1. Если есть рядом с .exe — портабельный.
    2. Иначе — классический.
    """
    portable = get_portable_config_path()
    if portable.exists():
        return portable
    return get_classic_config_path()


# ---------------------------------------------------------------------------
# Загрузка / сохранение
# ---------------------------------------------------------------------------

def load_config() -> Dict[str, Any]:
    """Загружает config.json. При ошибке — пустой словарь."""
    path = get_config_path()
    if not path.exists():
        return {}
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            return {}
        return data
    except (json.JSONDecodeError, OSError):
        return {}


def save_config(config: Dict[str, Any]) -> None:
    """Сохраняет config.json атомарно."""
    path = get_config_path()
    path.parent.mkdir(parents=True, exist_ok=True)

    tmp = path.with_suffix(".json.tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)

    tmp.replace(path)


def get(key: str, default: Any = None) -> Any:
    return load_config().get(key, default)


def update(**kwargs) -> None:
    cfg = load_config()
    cfg.update(kwargs)
    save_config(cfg)


# ---------------------------------------------------------------------------
# Маркеры
# ---------------------------------------------------------------------------

def get_markers() -> List[Dict[str, Any]]:
    """Список маркеров. Если в конфиге нет — дефолтные."""
    markers = get("markers")
    if not markers or not isinstance(markers, list):
        return [dict(m) for m in DEFAULT_MARKERS]
    return markers


def save_markers(markers: List[Dict[str, Any]]) -> None:
    update(markers=markers)


# ---------------------------------------------------------------------------
# Настройки (структура папок, поведение)
# ---------------------------------------------------------------------------

def get_settings() -> Dict[str, Any]:
    """Все настройки с дефолтными значениями."""
    cfg = load_config()
    result = dict(DEFAULT_SETTINGS)
    for k, v in DEFAULT_SETTINGS.items():
        if k in cfg:
            result[k] = cfg[k]
    return result


def save_settings(settings: Dict[str, Any]) -> None:
    update(**settings)


# ---------------------------------------------------------------------------
# Разрешение относительных путей (портабельный режим)
# ---------------------------------------------------------------------------

def resolve_path(relative: str) -> Path:
    """Раскрывает относительный путь от папки установки.

    Если путь абсолютный — возвращается как есть.
    Если относительный — <install_dir>/<relative>.
    Пустая строка — возвращает папку установки.
    """
    if not relative:
        return get_install_dir()
    p = Path(relative)
    if p.is_absolute():
        return p
    return get_install_dir() / p


def to_relative(path: str) -> str:
    """Преобразует абсолютный путь в относительный, если он внутри install_dir.

    Иначе возвращает путь как есть.
    """
    if not path:
        return ""
    p = Path(path).resolve()
    install = get_install_dir().resolve()
    try:
        return str(p.relative_to(install))
    except ValueError:
        return str(p)


# ---------------------------------------------------------------------------
# Портабельная структура — автосоздание при первом запуске
# ---------------------------------------------------------------------------

def ensure_portable_structure() -> Dict[str, str]:
    """Создаёт структуру папок и файлов при первом запуске.

    Возвращает словарь с относительными путями:
        {"excel": "clients.xlsx", "templates": "Шаблоны", "output": "Готовые документы"}
    """
    install = get_install_dir()
    install.mkdir(parents=True, exist_ok=True)

    paths = {
        "excel": "clients.xlsx",
        "templates": "Шаблоны",
        "output": "Готовые документы",
    }

    # 1. Создаём папки
    (install / paths["templates"]).mkdir(parents=True, exist_ok=True)
    (install / paths["output"]).mkdir(parents=True, exist_ok=True)

    # 2. Пример clients.xlsx (если нет)
    xlsx = install / paths["excel"]
    if not xlsx.exists():
        _create_example_xlsx(xlsx)

    # 3. Пример шаблона (если нет)
    tpl = install / paths["templates"] / "Пример_Договор.docx"
    if not tpl.exists():
        _create_example_template(tpl)

    # 4. config.json (если нет или он пустой/без relative_paths)
    portable = get_portable_config_path()
    classic = get_classic_config_path()

    need_create = True
    if portable.exists():
        try:
            existing = json.loads(portable.read_text(encoding="utf-8"))
            if isinstance(existing, dict) and "relative_paths" in existing:
                need_create = False
        except (json.JSONDecodeError, OSError):
            need_create = True

    if need_create:
        # Попробуем миграцию из классического
        migrated = False
        if classic.exists():
            try:
                with classic.open("r", encoding="utf-8") as f:
                    old_cfg = json.load(f)
                if isinstance(old_cfg, dict) and old_cfg:
                    # 1. Пути переносим в relative_paths
                    old_cfg.setdefault("relative_paths", {})
                    for key in ("excel_path", "templates_dir", "output_dir"):
                        if key in old_cfg:
                            old_cfg["relative_paths"][key] = old_cfg.pop(key)

                    # 2. Дополняем дефолтными маркерами, если их нет
                    if not old_cfg.get("markers"):
                        old_cfg["markers"] = [dict(m) for m in DEFAULT_MARKERS]

                    # 3. Дополняем дефолтными настройками
                    for k, v in DEFAULT_SETTINGS.items():
                        old_cfg.setdefault(k, v)

                    # 4. Убираем устаревший ключ
                    old_cfg.pop("selected_formats", None)

                    save_config_at(portable, old_cfg)
                    migrated = True
            except (json.JSONDecodeError, OSError):
                pass

        if not migrated:
            # Свежий конфиг с относительными путями
            default_cfg = {
                "relative_paths": {
                    "excel_path": paths["excel"],
                    "templates_dir": paths["templates"],
                    "output_dir": paths["output"],
                },
                "markers": [dict(m) for m in DEFAULT_MARKERS],
                **DEFAULT_SETTINGS,
            }
            save_config_at(portable, default_cfg)

    return paths


def save_config_at(path: Path, config: Dict[str, Any]) -> None:
    """Сохраняет config.json по конкретному пути (для миграции)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".json.tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
    tmp.replace(path)


# ---------------------------------------------------------------------------
# Генерация примеров — clients.xlsx и шаблона DOCX
# ---------------------------------------------------------------------------

def _create_example_xlsx(path: Path) -> None:
    """Создаёт пример clients.xlsx с 3 строками."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Клиенты"

    headers = ["ID", "ФИО", "Сумма", "Дата", "НомерДоговора", "Адрес", "Телефон"]
    ws.append(headers)

    rows = [
        # 1. Полностью заполненная
        [1, "Иванов Иван Иванович", 150000, "01.01.2024", "Д-001",
         "г. Москва, ул. Ленина, д. 1", "+7 900 123-45-67"],
        # 2. Частично заполненная (без адреса и телефона)
        [2, "Петрова Мария Сергеевна", 230500, "15.02.2024", "Д-002", "", ""],
        # 3. Пустой ID — при чтении пропускается
        ["", "Сидоров Пётр Алексеевич", 98750, "20.03.2024", "Д-003",
         "г. Казань", "+7 900 555-11-22"],
    ]

    for r in rows:
        ws.append(r)

    # Ширина колонок — поудобнее
    widths = [8, 30, 12, 12, 16, 40, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(path)
    wb.close()


def _create_example_template(path: Path) -> None:
    """Создаёт пример шаблона DOCX с маркерами."""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Заголовок
    h = doc.add_heading("ДОГОВОР № {{ContractNum}}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("от {{ContractDate}}")

    doc.add_paragraph("")
    doc.add_paragraph(
        "Гражданин(ка) {{ClientName}}, "
        "идентификатор {{ID}}, "
        "проживающий(ая) по адресу: {{Address}}, "
        "заключает настоящий договор на сумму {{Amount}} рублей."
    )

    doc.add_paragraph("")
    doc.add_paragraph("Телефон для связи: {{Phone}}.")
    doc.add_paragraph("")
    doc.add_paragraph("Подпись сторон: ___________________________")

    # Таблица с реквизитами
    doc.add_paragraph("")
    doc.add_paragraph("Реквизиты:")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"

    cells = [
        ("Поле", "Значение"),
        ("ФИО", "{{ClientName}}"),
        ("Сумма", "{{Amount}} руб."),
        ("Договор", "№ {{ContractNum}} от {{ContractDate}}"),
    ]
    for i, (k, v) in enumerate(cells):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v

    # Колонтитулы
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Договор № {{ContractNum}} — {{ClientName}}"
    section.footer.paragraphs[0].text = "Тел.: {{Phone}} | ID: {{ID}}"

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


# ---------------------------------------------------------------------------
# Демонстрация (для проверки)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    print("Платформа:", sys.platform)
    print("Папка установки:", get_install_dir())
    print("Портабельный режим:", is_portable())
    print("Активный config:", get_config_path())
    print()
    print("Маркеры:")
    for m in get_markers():
        print(f"  {m['marker']:<15} ← {m['column']:<20} ({m['format']})")
    print()
    print("Настройки:", get_settings())
