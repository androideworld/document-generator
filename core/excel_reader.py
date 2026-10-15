"""Чтение клиентов из clients.xlsx через openpyxl.

Чтение ведётся по имени колонки из шапки (первая строка).
Это позволяет:
    - менять порядок колонок,
    - добавлять свои колонки,
    - использовать любые имена в шапке.
"""

import re
from dataclasses import dataclass, asdict, field
from datetime import datetime, date
from typing import List, Dict
import openpyxl


@dataclass
class Client:
    """Клиент. Основные поля заполняются из «известных» колонок,
    raw — словарь со всеми колонками (по имени из шапки)."""
    id: str = ""
    fio: str = ""
    sum: str = ""
    date: str = ""
    contract_num: str = ""
    address: str = ""
    phone: str = ""
    raw: Dict[str, str] = field(default_factory=dict)

    def to_dict(self) -> dict:
        return asdict(self)


# Соответствие «известное поле Client → имя колонки по умолчанию».
# Используется, если в шапке есть эти колонки.
DEFAULT_COLUMNS = {
    "id": "ID",
    "fio": "ФИО",
    "sum": "Сумма",
    "date": "Дата",
    "contract_num": "НомерДоговора",
    "address": "Адрес",
    "phone": "Телефон",
}


# ---------------------------------------------------------------------------
# Нормализация значений
# ---------------------------------------------------------------------------

def _normalize_id(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).strftime("%d.%m.%Y")
        except ValueError:
            continue
    return text


def _cell(row: tuple, index: int):
    if index is None or index < 0:
        return None
    if index < len(row):
        return row[index]
    return None


# ---------------------------------------------------------------------------
# Чтение
# ---------------------------------------------------------------------------

def get_excel_columns(path: str) -> List[str]:
    """Возвращает список имён колонок из шапки (первая строка)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), tuple())
    wb.close()

    columns: List[str] = []
    for v in first_row:
        columns.append(_normalize_text(v))
    return columns


def read_clients(path: str) -> List[Client]:
    """Читает клиентов из Excel.

    - Шапка — первая строка.
    - Значения читаются по имени колонки.
    - Строки с пустым ID пропускаются.
    - Все колонки попадают в client.raw.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        wb.close()
        return []

    # Карта {имя_колонки: индекс}
    header_names = [_normalize_text(v) for v in header]
    col_index = {name: i for i, name in enumerate(header_names) if name}

    clients: List[Client] = []

    for row in rows_iter:
        if row is None:
            continue

        # Собираем raw: {имя: значение}
        raw: Dict[str, str] = {}
        for name, idx in col_index.items():
            raw[name] = _normalize_text(_cell(row, idx))

        # Известные поля — по именам колонок
        id_val = _normalize_id(raw.get(DEFAULT_COLUMNS["id"], ""))
        if not id_val:
            continue  # пустой ID — пропускаем

        client = Client(
            id=id_val,
            fio=raw.get(DEFAULT_COLUMNS["fio"], ""),
            sum=_normalize_text(raw.get(DEFAULT_COLUMNS["sum"], "")),
            date=_normalize_date(raw.get(DEFAULT_COLUMNS["date"], "")),
            contract_num=raw.get(DEFAULT_COLUMNS["contract_num"], ""),
            address=raw.get(DEFAULT_COLUMNS["address"], ""),
            phone=raw.get(DEFAULT_COLUMNS["phone"], ""),
            raw=raw,
        )
        clients.append(client)

    wb.close()
    return clients


# ---------------------------------------------------------------------------
# Фильтрация
# ---------------------------------------------------------------------------

def filter_clients(
    clients: List[Client],
    fio: str = "",
    id: str = "",
    contract: str = "",
    phone: str = "",
) -> List[Client]:
    """Фильтрует клиентов по 4 полям без учёта регистра."""
    q_fio = fio.strip().lower()
    q_id = id.strip().lower()
    q_contract = contract.strip().lower()
    q_phone = phone.strip().lower()

    def phone_digits(value: str) -> str:
        return "".join(ch for ch in value if ch.isdigit())

    q_phone_digits = phone_digits(q_phone)

    result: List[Client] = []
    for c in clients:
        if q_fio and q_fio not in c.fio.lower():
            continue
        if q_id and q_id not in c.id.lower():
            continue
        if q_contract and q_contract not in c.contract_num.lower():
            continue
        if q_phone_digits and q_phone_digits not in phone_digits(c.phone):
            continue
        result.append(c)

    return result


# ---------------------------------------------------------------------------
# Утилиты для структуры папок
# ---------------------------------------------------------------------------

def short_fio(fio: str) -> str:
    """Иванов Иван Иванович → Иванов И.И."""
    parts = fio.strip().split()
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    surname = parts[0]
    initials = ".".join(p[0].upper() for p in parts[1:] if p) + "."
    return f"{surname} {initials}"


def sanitize_folder_name(name: str) -> str:
    """Убирает недопустимые символы из имени папки/файла."""
    cleaned = re.sub(r'[<>:"/\\|?*]', "_", name)
    cleaned = cleaned.rstrip(". ")
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned or "client"


# ---------------------------------------------------------------------------
# CLI для проверки
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "clients.xlsx"

    print(f"Файл: {path}")
    print(f"Колонки: {get_excel_columns(path)}")
    print()

    clients = read_clients(path)
    print(f"Прочитано клиентов: {len(clients)}")
    for i, c in enumerate(clients, 1):
        print(f"{i}. ID={c.id} | {c.fio} | дог. {c.contract_num} | {c.phone}")
        print(f"   raw = {c.raw}")
