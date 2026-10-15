"""Замена маркеров {{...}} в DOCX-шаблонах через python-docx.

Маркеры берутся из списка (обычно — из config.json через core.config.get_markers()).

Каждый маркер описывается словарём:
    {
        "marker": "ClientName",       # имя в шаблоне, без {{ }}
        "column": "ФИО",              # имя колонки в clients.xlsx
        "format": "text",             # text | amount | date | raw
        "required": True,             # обязателен ли (для предупреждений)
    }
"""

import re
from pathlib import Path
from typing import Iterable, List, Dict, Any, Optional, Set

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph

from .excel_reader import Client


# ---------------------------------------------------------------------------
# Форматирование значений
# ---------------------------------------------------------------------------

def format_amount(value: str) -> str:
    """Сумма с разделителем тысяч (пробел).

    '150000'      → '150 000'
    '150000.50'   → '150 000.50'
    '' / None     → ''
    """
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""

    m = re.match(r"^(\d+)([.,]\d+)?$", text.replace(" ", ""))
    if not m:
        return text

    integer_part, frac_part = m.group(1), m.group(2) or ""
    grouped = f"{int(integer_part):,}".replace(",", " ")
    return grouped + frac_part


def format_date(value: str) -> str:
    """Дата уже нормализована в excel_reader."""
    return (value or "").strip()


def format_value(value: str, fmt: str) -> str:
    """Форматирует значение по типу."""
    if fmt == "amount":
        return format_amount(value)
    if fmt == "date":
        return format_date(value)
    if fmt == "raw":
        return value or ""
    # text — по умолчанию
    return (value or "").strip()


# ---------------------------------------------------------------------------
# Построение карты маркеров
# ---------------------------------------------------------------------------

def build_marker_map(
    client: Client,
    markers: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, str]:
    """Строит карту {имя_маркера: значение}.

    Значения берутся из client.raw[column] и форматируются по marker["format"].

    Если значение пустое:
        - on_empty_value == "keep_marker" → маркер НЕ попадает в карту,
          останется в документе как есть;
        - on_empty_value == "placeholder" → подставляются прочерки
          (количество = placeholder_underscores);
        - on_empty_value == "empty_string" → подставляются пробелы
          (количество = empty_string_spaces).

    Если markers не заданы — используется get_markers() из core.config.
    """
    if markers is None:
        from .config import get_markers
        markers = get_markers()

    # Читаем настройки поведения
    try:
        from .config import get_settings
        settings = get_settings()
        on_empty = settings.get("on_empty_value", "empty_string")
        spaces_count = int(settings.get("empty_string_spaces", 5))
        underscores_count = int(settings.get("placeholder_underscores", 10))
    except Exception:
        on_empty = "empty_string"
        spaces_count = 5
        underscores_count = 10

    result: Dict[str, str] = {}
    for m in markers:
        marker_name = m.get("marker", "").strip()
        column = m.get("column", "").strip()
        fmt = m.get("format", "text")

        if not marker_name:
            continue

        raw_value = ""
        if column:
            raw_value = client.raw.get(column, "") if client.raw else ""

        raw_value = raw_value.strip() if isinstance(raw_value, str) else ""

        if not raw_value:
            if on_empty in ("keep_marker", "keep"):
                # Не добавляем маркер в карту — он останется в документе
                continue
            elif on_empty == "placeholder":
                result[marker_name] = "_" * max(0, underscores_count)
                continue
            else:  # empty_string
                result[marker_name] = " " * max(0, spaces_count)
                continue

        result[marker_name] = format_value(raw_value, fmt)

    return result

# ---------------------------------------------------------------------------
# Регулярка для маркеров
# ---------------------------------------------------------------------------

# Имя маркера: буквы (включая кириллицу), цифры, подчёркивания
MARKER_RE = re.compile(r"\{\{\s*([\w\u0400-\u04FF]+)\s*\}\}")


def replace_in_text(text: str, marker_map: Dict[str, str]) -> str:
    """Заменяет все {{...}} в строке. Неизвестные — оставляет как есть."""
    def repl(match: re.Match) -> str:
        key = match.group(1)
        if key in marker_map:
            return marker_map[key]
        return match.group(0)
    return MARKER_RE.sub(repl, text)


# ---------------------------------------------------------------------------
# Обход всех текстовых контейнеров
# ---------------------------------------------------------------------------

def _iter_paragraphs(container) -> Iterable[Paragraph]:
    for p in container.paragraphs:
        yield p
    for table in getattr(container, "tables", []):
        yield from _iter_table(table)


def _iter_table(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from _iter_paragraphs(cell)


def _replace_in_paragraph(paragraph: Paragraph, marker_map: Dict[str, str]) -> bool:
    """Заменяет маркеры в параграфе. Возвращает True, если что-то заменили."""
    changed = False

    # Быстрый путь: замена внутри каждого run
    for run in paragraph.runs:
        new_text = replace_in_text(run.text, marker_map)
        if new_text != run.text:
            run.text = new_text
            changed = True

    # Случай, когда Word разбил маркер на несколько run
    full = paragraph.text
    if "{{" in full and MARKER_RE.search(full):
        new_full = replace_in_text(full, marker_map)
        if new_full != full:
            for i, run in enumerate(paragraph.runs):
                if i == 0:
                    run.text = new_full
                else:
                    run.text = ""
            changed = True

    return changed


# ---------------------------------------------------------------------------
# Публичное API
# ---------------------------------------------------------------------------

def replace_markers(doc: DocumentType, client: Client,
                    markers: Optional[List[Dict[str, Any]]] = None) -> int:
    """Заменяет маркеры в документе (параграфы, таблицы, колонтитулы).

    Логика для НЕНАЙДЕННЫХ маркеров (есть в шаблоне, нет в конфиге)
    определяется настройкой on_missing_marker:
        - highlight_and_warn → оставить + подсветить + предупредить
        - highlight          → оставить + подсветить
        - keep               → оставить как есть
        - placeholder        → поставить прочерк (placeholder_underscores)
        - empty_string       → подставить пробелы (empty_string_spaces)
        - hide               → удалить маркер полностью
    """
    # Читаем настройки
    try:
        from .config import get_settings
        settings = get_settings()
        on_missing = settings.get("on_missing_marker", "highlight_and_warn")
        spaces_count = int(settings.get("empty_string_spaces", 5))
        underscores_count = int(settings.get("placeholder_underscores", 10))
    except Exception:
        on_missing = "highlight_and_warn"
        spaces_count = 5
        underscores_count = 10

    # 1. Карта для известных маркеров (из конфига)
    marker_map = build_marker_map(client, markers)

    # 2. Для подстановочных режимов — добавляем в карту ненайденные маркеры
    if on_missing in ("placeholder", "empty_string"):
        all_in_template = _collect_all_markers(doc)
        known_markers = set(marker_map.keys())
        unknown_markers = all_in_template - known_markers

        for name in unknown_markers:
            if on_missing == "placeholder":
                marker_map[name] = "_" * max(0, underscores_count)
            elif on_missing == "empty_string":
                marker_map[name] = " " * max(0, spaces_count)

    counter = 0

    for paragraph in _iter_paragraphs(doc):
        if _replace_in_paragraph(paragraph, marker_map):
            counter += 1

    for section in doc.sections:
        for header_footer in (section.header, section.footer):
            for paragraph in _iter_paragraphs(header_footer):
                if _replace_in_paragraph(paragraph, marker_map):
                    counter += 1

    # 3. Если режим hide — удаляем оставшиеся ненайденные маркеры
    if on_missing == "hide":
        _remove_unreplaced_markers(doc)

    return counter


def _collect_all_markers(doc: DocumentType) -> Set[str]:
    """Собирает все имена маркеров {{...}} в документе (без замены)."""
    found: Set[str] = set()

    def scan_text(text: str) -> None:
        for m in MARKER_RE.finditer(text):
            found.add(m.group(1))

    for p in _iter_paragraphs(doc):
        scan_text(p.text)
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in _iter_paragraphs(hf):
                scan_text(p.text)

    return found


def _remove_unreplaced_markers(doc: DocumentType) -> int:
    """Удаляет незаменённые маркеры {{...}} из документа полностью.

    Работает как с цельными, так и с разбитыми по runs маркерами.

    :return: количество удалённых маркеров
    """
    removed = 0

    def clean_paragraph(p: Paragraph) -> None:
        nonlocal removed
        if not MARKER_RE.search(p.text):
            return

        cleaned_text = MARKER_RE.sub("", p.text)
        removed += len(MARKER_RE.findall(p.text))

        if not p.runs:
            return
        first_run = p.runs[0]
        first_run.text = cleaned_text
        for run in list(p.runs[1:]):
            run.text = ""

    for p in _iter_paragraphs(doc):
        clean_paragraph(p)
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in _iter_paragraphs(hf):
                clean_paragraph(p)

    return removed


def list_markers_in_template(template_path: str) -> Set[str]:
    """Сканирует DOCX-шаблон и возвращает множество имён маркеров.

    Обходит параграфы, таблицы, колонтитулы.
    """
    doc = Document(str(template_path))
    found: Set[str] = set()

    def scan_text(text: str) -> None:
        for m in MARKER_RE.finditer(text):
            found.add(m.group(1))

    def scan_container(container) -> None:
        for p in container.paragraphs:
            scan_text(p.text)
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    scan_container(cell)

    scan_container(doc)
    for section in doc.sections:
        scan_container(section.header)
        scan_container(section.footer)

    return found


def find_unreplaced_markers(doc: DocumentType) -> Set[str]:
    """Возвращает множество маркеров, оставшихся в документе после замены."""
    found: Set[str] = set()

    def scan_text(text: str) -> None:
        for m in MARKER_RE.finditer(text):
            found.add(m.group(1))

    for p in _iter_paragraphs(doc):
        scan_text(p.text)
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in _iter_paragraphs(hf):
                scan_text(p.text)

    return found

def highlight_unreplaced_markers(doc: DocumentType, color: str = "red") -> int:
    """Подсвечивает незаменённые маркеры {{...}} в документе.

    Всегда подсвечивается ТОЛЬКО маркер, а не весь run или параграф.
    Работает с любым случаем:
        - один run с текстом + маркером,
        - маркер, разбитый на несколько runs,
        - несколько маркеров в одном параграфе.

    :param doc: документ после replace_markers
    :param color: 'red' (по умолчанию)
    :return: количество подсвеченных маркеров
    """
    from docx.enum.text import WD_COLOR_INDEX

    colors = {
        "red": WD_COLOR_INDEX.RED,
        "yellow": WD_COLOR_INDEX.YELLOW,
        "green": WD_COLOR_INDEX.BRIGHT_GREEN,
        "bright_green": WD_COLOR_INDEX.BRIGHT_GREEN,
        "turquoise": WD_COLOR_INDEX.TURQUOISE,
        "pink": WD_COLOR_INDEX.PINK,
    }
    hl_color = colors.get(color.lower(), WD_COLOR_INDEX.RED)

    count = 0

    def process_paragraph(p: Paragraph) -> None:
        nonlocal count

        full_text = p.text
        if not MARKER_RE.search(full_text):
            return

        if not p.runs:
            return

        # --- Шаг 1. Склеиваем все runs в один, чтобы иметь цельный текст ---
        # Запоминаем формат первого run как базовый
        first_run = p.runs[0]
        # Сохраняем копии форматирования
        font_name = first_run.font.name
        font_size = first_run.font.size
        font_bold = first_run.font.bold
        font_italic = first_run.font.italic
        font_underline = first_run.font.underline

        # Удаляем все runs, кроме первого
        for run in list(p.runs[1:]):
            run._element.getparent().remove(run._element)

        # Очищаем первый run
        first_run.text = ""
        first_run.font.highlight_color = None

        # --- Шаг 2. Разбиваем full_text на части ---
        parts = []
        pos = 0
        for m in MARKER_RE.finditer(full_text):
            if m.start() > pos:
                parts.append(("text", full_text[pos:m.start()]))
            parts.append(("marker", m.group(0)))
            pos = m.end()
        if pos < len(full_text):
            parts.append(("text", full_text[pos:]))

        # --- Шаг 3. Пересоздаём runs ---
        # Первый run используем как «базовый» — если первой части нет
        # (текст начинается сразу с маркера), оставим его пустым
        first_used = False
        for kind, text in parts:
            if not first_used:
                run = first_run
                run.text = text
                first_used = True
            else:
                run = p.add_run(text)

            # Восстанавливаем форматирование
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = font_bold
            run.font.italic = font_italic
            run.font.underline = font_underline

            if kind == "marker":
                run.font.highlight_color = hl_color
                count += 1

    for p in _iter_paragraphs(doc):
        process_paragraph(p)
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in _iter_paragraphs(hf):
                process_paragraph(p)

    return count

def generate_docx(
    template_path: str,
    output_path: str,
    client: Client,
    markers: Optional[List[Dict[str, Any]]] = None,
    return_unreplaced: bool = False,
):
    """Открывает шаблон, заменяет маркеры, сохраняет результат.

    Если в настройках включена подсветка (on_missing_marker = "highlight"
    или "highlight_and_warn"), незаменённые маркеры подсвечиваются
    красным в самом документе.

    :param return_unreplaced: если True — вернуть set незаменённых маркеров
    """
    tpl = Path(template_path)
    if not tpl.exists():
        raise FileNotFoundError(f"Шаблон не найден: {template_path}")

    doc = Document(str(tpl))
    replace_markers(doc, client, markers)

    # Читаем настройку подсветки
    try:
        from .config import get_settings
        on_missing = get_settings().get("on_missing_marker", "highlight_and_warn")
    except Exception:
        on_missing = "highlight_and_warn"

    # Находим незаменённые ДО сохранения
    unreplaced = find_unreplaced_markers(doc)

    # Подсвечиваем, если нужно
    if unreplaced and on_missing in ("highlight", "highlight_and_warn"):
        highlight_unreplaced_markers(doc, color="red")

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    if return_unreplaced:
        return unreplaced
    return None


# ---------------------------------------------------------------------------
# CLI для проверки
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Использование: python core/docx_generator.py <template.docx>")
        sys.exit(0)

    tpl = sys.argv[1]
    print(f"Шаблон: {tpl}")
    found = list_markers_in_template(tpl)
    print(f"Маркеры в шаблоне ({len(found)}):")
    for m in sorted(found):
        print(f"  {{{{ {m} }}}}")
